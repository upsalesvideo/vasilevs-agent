# -*- coding: utf-8 -*-
"""
build_script.py — собирает «презентацию текстом»: продающий текст копирайтера к каждому
слайду (поле `copy` в project.json) + заголовок и текст слайда, подряд, как один рассказ.

Пишет два файла в out_dir:
  <project_name>_текст.md    — Markdown (заголовок слайда, текст на слайде, продающий текст)
  <project_name>_текст.docx  — то же в Word (нужен python-docx)

Запуск:  python build_script.py project.json [--no-docx]
"""
import os
import sys
import json

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


def safe_name(name):
    name = (name or "presentation").replace("/", "_").replace("\\", "_")
    for ch in ':*?"<>|':
        name = name.replace(ch, "_")
    return name


def build_md(cfg, slides):
    lines = ["# " + (cfg.get("project_name") or "Презентация"), ""]
    for s in slides:
        lines.append("## %d. %s" % (s["num"], s.get("title", "")))
        lines.append("")
        if s.get("text"):
            lines.append("*На слайде:* " + s["text"])
            lines.append("")
        if s.get("copy"):
            lines.append(s["copy"].strip())
            lines.append("")
        else:
            lines.append("_(текст копирайтера не написан)_")
            lines.append("")
    return "\n".join(lines)


def build_docx(cfg, slides, out_path):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError:
        print("python-docx не установлен — .docx пропущен (pip install python-docx)")
        return False
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)
    doc.add_heading(cfg.get("project_name") or "Презентация", level=0)
    for s in slides:
        doc.add_heading("%d. %s" % (s["num"], s.get("title", "")), level=1)
        if s.get("text"):
            p = doc.add_paragraph()
            r = p.add_run("На слайде: ")
            r.bold = True
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            r2 = p.add_run(s["text"])
            r2.italic = True
            r2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        body = (s.get("copy") or "").strip()
        if body:
            for para in body.split("\n\n"):
                para = para.strip()
                if para:
                    doc.add_paragraph(para.replace("\n", " "))
        else:
            doc.add_paragraph("(текст копирайтера не написан)")
    doc.save(out_path)
    return True


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    cfg_path = sys.argv[1]
    with open(cfg_path, "r", encoding="utf-8") as f:
        cfg = json.load(f)
    cfg_dir = os.path.dirname(os.path.abspath(cfg_path))
    out_dir = cfg["out_dir"]
    if not os.path.isabs(out_dir):
        out_dir = os.path.join(cfg_dir, out_dir)
    os.makedirs(out_dir, exist_ok=True)
    slides = sorted(cfg.get("slides", []), key=lambda s: s["num"])
    base = os.path.join(out_dir, safe_name(cfg.get("project_name")) + "_текст")

    md = build_md(cfg, slides)
    with open(base + ".md", "w", encoding="utf-8") as f:
        f.write(md)
    words = sum(len((s.get("copy") or "").split()) for s in slides)
    missing = [s["num"] for s in slides if not (s.get("copy") or "").strip()]
    print("MD:   %s (%d слайдов, %d слов текста копирайтера)" % (base + ".md", len(slides), words))
    if missing:
        print("   !! без текста копирайтера: слайды %s" % ", ".join(map(str, missing)))
    if "--no-docx" not in sys.argv and build_docx(cfg, slides, base + ".docx"):
        print("DOCX: %s" % (base + ".docx"))
    print("OK")


if __name__ == "__main__":
    main()
